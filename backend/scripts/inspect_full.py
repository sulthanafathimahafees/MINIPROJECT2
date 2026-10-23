import docx

doc = docx.Document(r'C:\Users\ASUS\Downloads\ScholarNexus_Report.docx')

print('=== HEADER / FOOTER ANALYSIS ===')
sample_secs = [0, 1, 2, 5, 10, 20, 30, 50, 70, 80]
for i in sample_secs:
    if i < len(doc.sections):
        sec = doc.sections[i]
        hdr = sec.header
        ftr = sec.footer
        hdr_text = ' '.join(p.text.strip() for p in hdr.paragraphs if p.text.strip())
        ftr_text = ' '.join(p.text.strip() for p in ftr.paragraphs if p.text.strip())
        print(f'Sec {i:2d} | Hdr: "{hdr_text}" | Ftr: "{ftr_text}"')

print('\n=== FIGURE CAPTIONS ===')
fig_count = 0
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if ('fig' in t.lower() or 'figure' in t.lower()) and len(t) < 100:
        align = str(p.alignment)
        runs_desc = [(r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs[:2]]
        print(f'P{i:3d} [{p.style.name:12}] align={align:10} runs={runs_desc} : {t}')
        fig_count += 1
        if fig_count >= 15:
            break

print('\n=== TABLE CAPTIONS ===')
tbl_count = 0
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if ('table' in t.lower()) and len(t) < 100:
        align = str(p.alignment)
        runs_desc = [(r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs[:2]]
        print(f'P{i:3d} [{p.style.name:12}] align={align:10} runs={runs_desc} : {t}')
        tbl_count += 1
        if tbl_count >= 15:
            break

print('\n=== CHAPTER 4 (SYSTEM DESIGN) STRUCTURE ===')
for i in range(450, 540):
    p = doc.paragraphs[i]
    t = p.text.strip()
    if p.style.name.startswith('Heading') and t:
        print(f'P{i:3d} [{p.style.name:12}] align={p.alignment} : {t}')
